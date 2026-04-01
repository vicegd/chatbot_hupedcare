import os, re, base64, json, yaml
from datetime import datetime
from bs4 import BeautifulSoup
import PyPDF2
from docx import Document
from moviepy import VideoFileClip
from openai import OpenAI
import subprocess
from docx import Document

#Initialization
def load_config():
    with open("./config/config.yaml", "r") as f:
        return yaml.safe_load(f)
    
def load_metadata():
    METADATA_FILE = os.path.join(config['storage']['data_folder'], ".metadata.json")
    if os.path.exists(METADATA_FILE):
        with open(METADATA_FILE, "r") as f:
            return json.load(f)
    return {}

config = load_config()
metadata = load_metadata()
model = OpenAI(base_url=config['ai']['base_url'], api_key=os.getenv("MODEL_API_KEY"))

def extract_clean_text(text):
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    return "\n".join(lines)

def extract_from_html_or_php(file_path):
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        content = f.read()
    content = re.sub(r'<\?php.*?\?>', '', content, flags=re.DOTALL | re.IGNORECASE)
    soup = BeautifulSoup(content, 'html.parser')
    for tag in soup(["script", "style", "header", "footer", "nav", "aside"]):
        tag.decompose()
    return extract_clean_text(soup.get_text(separator=' '))

def extract_from_pdf(file_path):
    text = ""
    try:
        with open(file_path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                text += page.extract_text() + "\n"
    except Exception as e: print(f"PDF Error: {e}")
    return extract_clean_text(text)

def extract_description_from_image_with_ai(image_path):
    try:
        with open(image_path, "rb") as img:
            b64_img = base64.b64encode(img.read()).decode('utf-8')
        response = model.chat.completions.create(
            model=config['ai']['vision_model'],
            messages=[{"role": "user", "content": [
                {"type": "text", "text": "Describe this image in detail for a RAG system. Transcribe any text."},
                {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{b64_img}"}}
            ]}],
            max_tokens=300
        )
        return response.choices[0].message.content
    except Exception as e: return f"Vision Error: {e}"

def extract_description_from_video_with_ai(video_path):
    try:
        clip = VideoFileClip(video_path)
        audio_path = video_path + ".mp3"
        clip.audio.write_audiofile(audio_path, logger=None)
        with open(audio_path, "rb") as f:
            transcription = model.audio.transcriptions.create(model="whisper-large-v3", file=f).text
        
        frame_path = video_path + "_snap.jpg"
        clip.save_frame(frame_path, t=clip.duration/2)
        visual_desc = extract_description_from_image_with_ai(frame_path)
        
        os.remove(audio_path); os.remove(frame_path)
        return f"VIDEO VISUAL: {visual_desc}\nTRANSCRIPTION: {transcription}"
    except Exception as e: return f"Video Error: {e}"

def extract_from_docx(file_path):
    """Extrae texto de archivos .docx modernos."""
    try:
        doc = Document(file_path)
        full_text = []
        for para in doc.paragraphs:
            full_text.append(para.text)
        # También extraemos texto de tablas, que a veces se olvida
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    full_text.append(cell.text)
        return extract_clean_text("\n".join(full_text))
    except Exception as e:
        print(f"Error procesando DOCX {file_path}: {e}")
        return ""

def extract_from_doc(file_path):
    """
    Extrae texto de archivos .doc (antiguos).
    Nota: Requiere que el sistema tenga instalado 'antiword'.
    Si estás en Windows, lo ideal es convertirlos a .docx primero.
    """
    try:
        # Intentamos usar antiword (comando de consola muy ligero)
        result = subprocess.run(['antiword', file_path], capture_output=True, text=True, encoding='utf-8', errors='ignore')
        if result.returncode == 0:
            return extract_clean_text(result.stdout)
        else:
            return "Error: .doc format requires 'antiword' installed on the server."
    except FileNotFoundError:
        return "Error: System tool 'antiword' not found for .doc processing."
    except Exception as e:
        return f"Error procesando DOC (legacy): {e}"